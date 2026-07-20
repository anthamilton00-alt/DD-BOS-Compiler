from pathlib import Path
import sys

from docx import Document


PROJECT_ROOT = Path(__file__).resolve().parents[1]
SRC_FOLDER = PROJECT_ROOT / "src"

if str(SRC_FOLDER) not in sys.path:
    sys.path.insert(0, str(SRC_FOLDER))

from core.metadata import extract_metadata


def create_test_document(file_path: Path, paragraphs: list[str]) -> None:
    document = Document()

    for text in paragraphs:
        document.add_paragraph(text)

    document.save(file_path)


def test_standard_assembly_metadata(tmp_path: Path) -> None:
    file_path = tmp_path / "ASM-202 — Dual Cat6 Home Run Assembly.docx"

    create_test_document(
        file_path,
        [
            "DIGITALLY DEFINED",
            "ASSEMBLY ID:",
            "ASM-202",
            "ASSEMBLY NAME:",
            "Dual Cat6 Home Run Assembly",
            "VERSION:",
            "1.0",
            "STATUS:",
            "FROZEN",
        ],
    )

    metadata = extract_metadata(file_path)

    assert metadata["Document ID"] == "ASM-202"
    assert metadata["Title"] == "Dual Cat6 Home Run Assembly"
    assert metadata["Version"] == "1.0"
    assert metadata["Status"] == "FROZEN"


def test_filename_supplies_document_title(tmp_path: Path) -> None:
    file_path = tmp_path / "DD-007 — Builder Package Library.docx"

    create_test_document(
        file_path,
        [
            "DIGITALLY DEFINED",
            "DOCUMENT 007",
            "BUILDER PACKAGE LIBRARY",
            "DOCUMENT ID:",
            "DD-007",
            "VERSION:",
            "1.0",
            "STATUS:",
            "FROZEN",
            "DOCUMENT OWNER:",
            "Digitally Defined",
        ],
    )

    metadata = extract_metadata(file_path)

    assert metadata["Document ID"] == "DD-007"
    assert metadata["Title"] == "Builder Package Library"
    assert metadata["Version"] == "1.0"
    assert metadata["Status"] == "FROZEN"
    assert metadata["Owner"] == "Digitally Defined"


def test_floor_plan_placeholders_are_rejected(tmp_path: Path) -> None:
    file_path = tmp_path / "FP-0001 — Master Floor Plan Template.docx"

    create_test_document(
        file_path,
        [
            "FLOOR PLAN ID:",
            "FP-0001",
            "FLOOR PLAN NAME:",
            "Builder Floor Plan Number:",
            "VERSION:",
            "Date",
            "STATUS:",
            "Date Approved:",
        ],
    )

    metadata = extract_metadata(file_path)

    assert metadata["Document ID"] == "FP-0001"
    assert metadata["Title"] == "Master Floor Plan Template"
    assert metadata["Version"] == ""
    assert metadata["Status"] == ""


def test_room_title_comes_from_filename(tmp_path: Path) -> None:
    file_path = tmp_path / "ROOM-001 — Great Room Standard.docx"

    create_test_document(
        file_path,
        [
            "ROOM ID:",
            "ROOM-001",
            "ROOM NAME:",
            "Great Room",
            "VERSION:",
            "1.0",
            "STATUS:",
            "FROZEN",
        ],
    )

    metadata = extract_metadata(file_path)

    assert metadata["Document ID"] == "ROOM-001"
    assert metadata["Title"] == "Great Room Standard"
    assert metadata["Version"] == "1.0"
    assert metadata["Status"] == "FROZEN"